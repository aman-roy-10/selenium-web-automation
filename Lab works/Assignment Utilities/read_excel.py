from openpyxl import load_workbook

# Load Excel file
workbook = load_workbook("testdata.xlsx")

# Select Sheet1
sheet = workbook["Sheet1"]

# Read rows
for row in sheet.iter_rows(values_only=True):
    print(row)

workbook.close()