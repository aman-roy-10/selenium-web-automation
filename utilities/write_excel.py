from openpyxl import load_workbook

# Load existing Excel file
workbook = load_workbook("testdata.xlsx")

# Select Sheet1
sheet = workbook["Sheet1"]

# Find the next empty row
next_row = sheet.max_row + 1

# Write new data
sheet.cell(row=next_row, column=1).value = "Neha Singh"
sheet.cell(row=next_row, column=2).value = "neha@gmail.com"

# Save the Excel file
workbook.save("testdata.xlsx")

print("Data written successfully")
print("Added:", sheet.cell(row=next_row, column=1).value)
print("Email:", sheet.cell(row=next_row, column=2).value)

workbook.close()